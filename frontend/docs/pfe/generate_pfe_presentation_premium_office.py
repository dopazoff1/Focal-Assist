from __future__ import annotations

from pathlib import Path
import math
import shutil

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "pfe"
SHOT_DIR = OUT_DIR / "screenshots-full"
LOGO = ROOT / "tmp_template_image1.png"
OUT = OUT_DIR / "presentation-pfe-focal-assist-v3-premium.pptx"
OFFICE_COPY = OUT_DIR / "presentation-pfe-focal-assist-v3-premium-office-native.pptx"


SLIDE_W = 13.333333
SLIDE_H = 7.5
PX_SCALE = 72.0

C = {
    "navy": "091430",
    "ink": "182235",
    "muted": "66758D",
    "blue": "2E63EB",
    "cyan": "00A7A5",
    "green": "16A34A",
    "orange": "EA580C",
    "red": "DC2626",
    "bg": "F7FAFF",
    "white": "FFFFFF",
    "border": "CAD8EA",
    "soft_blue": "EAF1FF",
    "soft_cyan": "E7FAF7",
    "dark": "111827",
}


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.strip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def u(px: float) -> float:
    return px / PX_SCALE


def add_shape(slide, shape, x, y, w, h, fill=C["white"], line=C["border"], radius=True):
    shp = slide.shapes.add_shape(shape, Inches(u(x)), Inches(u(y)), Inches(u(w)), Inches(u(h)))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(0.65)
    else:
        shp.line.fill.background()
    return shp


def add_rect(slide, x, y, w, h, fill=C["white"], line=C["border"], round_rect=True):
    return add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE if round_rect else MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line)


def add_line(slide, x1, y1, x2, y2, color=C["border"], width=1.0):
    line = slide.shapes.add_connector(1, Inches(u(x1)), Inches(u(y1)), Inches(u(x2)), Inches(u(y2)))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def estimate_single_line_width(text: str, size: float, bold: bool = False) -> float:
    weights = {
        "i": 0.25,
        "l": 0.28,
        "I": 0.30,
        " ": 0.32,
        ".": 0.24,
        ",": 0.24,
        ":": 0.24,
        ";": 0.24,
        "-": 0.32,
        "|": 0.22,
        "W": 0.86,
        "M": 0.82,
        "w": 0.72,
        "m": 0.78,
    }
    total = 0.0
    for ch in text:
        total += weights.get(ch, 0.56 if ch.islower() else 0.62)
    return total * size * (1.06 if bold else 1.0) + 8


def text_box(slide, text, x, y, w, h, size=14, color=C["ink"], bold=False, align="left", fill=None, tight=None):
    lines = str(text).split("\n") or [""]
    if tight is None:
        tight = fill is None and len(lines) == 1 and len(lines[0]) <= 62 and h <= 48
    if tight and len(lines) == 1:
        estimated_w = min(w, max(8, estimate_single_line_width(lines[0], size, bold)))
        estimated_h = min(h, max(7, size * 1.42))
        if align == "right":
            x += w - estimated_w
        elif align == "center":
            x += (w - estimated_w) / 2
        w = estimated_w
        h = estimated_h
    box = slide.shapes.add_textbox(Inches(u(x)), Inches(u(y)), Inches(u(w)), Inches(u(h)))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    tf.word_wrap = not tight
    tf.vertical_anchor = MSO_ANCHOR.TOP
    if tight:
        tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
        if align == "center":
            p.alignment = PP_ALIGN.CENTER
        elif align == "right":
            p.alignment = PP_ALIGN.RIGHT
        else:
            p.alignment = PP_ALIGN.LEFT
    if fill:
        box.fill.solid()
        box.fill.fore_color.rgb = rgb(fill)
    return box


def bullet_box(slide, items, x, y, w, h, size=12, color=C["ink"]):
    box = slide.shapes.add_textbox(Inches(u(x)), Inches(u(y)), Inches(u(w)), Inches(u(h)))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.level = 0
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(5)
    return box


def background(slide, section="PFE", number=1):
    add_rect(slide, 0, 0, 960, 540, C["bg"], None, False)
    # Soft technical atmosphere.
    add_shape(slide, MSO_SHAPE.OVAL, 704, 20, 245, 245, C["soft_cyan"], None)
    add_shape(slide, MSO_SHAPE.OVAL, -98, 112, 250, 250, C["soft_blue"], None)
    for gx in range(0, 961, 96):
        add_line(slide, gx, 58, gx, 540, "E2ECF7", 0.25)
    for gy in range(64, 541, 72):
        add_line(slide, 0, gy, 960, gy, "E2ECF7", 0.25)
    add_rect(slide, 0, 0, 960, 48, C["navy"], None, False)
    add_rect(slide, 0, 46, 960, 2.2, C["cyan"], None, False)
    add_rect(slide, 24, 13, 24, 24, C["blue"], C["cyan"])
    text_box(slide, "KA", 29, 18, 18, 10, 8, C["white"], True)
    text_box(slide, "Focal V3", 58, 12, 145, 14, 10, C["white"], True)
    text_box(slide, "PFE | ISMAGI", 58, 27, 120, 10, 7, "B0DCEB")
    text_box(slide, section.upper(), 690, 14, 190, 14, 8, "B0DCEB", True, "right")
    text_box(slide, f"{number:02}", 902, 13, 34, 18, 10, C["white"], True)


def section_title(slide, label, title, subtitle=""):
    text_box(slide, label.upper(), 54, 61, 260, 16, 8, C["cyan"], True)
    text_box(slide, title, 54, 84, 845, 38, 25, C["navy"], True)
    add_rect(slide, 54, 132, 92, 4, C["cyan"], None, False)
    add_rect(slide, 154, 132, 42, 4, C["blue"], None, False)
    if subtitle:
        text_box(slide, subtitle, 54, 148, 820, 28, 10.5, C["muted"])


def add_card(slide, title, body, x, y, w=250, h=116, accent=C["blue"]):
    add_rect(slide, x, y, w, h, C["white"], C["border"])
    add_rect(slide, x, y, w, 5, accent, None, False)
    text_box(slide, title, x + 18, y + 16, w - 36, 20, 14, accent, True)
    text_box(slide, body, x + 18, y + 47, w - 36, h - 56, 9, C["ink"])


def panel(slide, title, items, x=78, y=176, w=805, h=230, accent=C["blue"]):
    add_rect(slide, x, y, w, h, C["white"], C["border"])
    add_rect(slide, x, y, w, 5, accent, None, False)
    text_box(slide, title, x + 24, y + 20, w - 48, 22, 16, accent, True)
    bullet_box(slide, items, x + 28, y + 58, w - 56, h - 74, 11)


def fit_image(path: Path, x, y, w, h):
    with Image.open(path) as img:
        aspect = img.width / img.height
    box_aspect = w / h
    if aspect > box_aspect:
        nw = w
        nh = w / aspect
        return x, y + (h - nh) / 2, nw, nh
    nh = h
    nw = h * aspect
    return x + (w - nw) / 2, y, nw, nh


def add_picture_fit(slide, path: Path, x, y, w, h):
    if not path.exists():
        add_rect(slide, x, y, w, h, C["white"], C["red"])
        text_box(slide, f"Image introuvable\n{path.name}", x + 20, y + 20, w - 40, h - 40, 12, C["red"], True)
        return
    px, py, pw, ph = fit_image(path, x, y, w, h)
    slide.shapes.add_picture(str(path), Inches(u(px)), Inches(u(py)), Inches(u(pw)), Inches(u(ph)))


def screenshot(slide, path: Path, x, y, w, h):
    add_rect(slide, x, y, w, h, C["navy"], C["border"])
    add_rect(slide, x + 8, y + 8, w - 16, 26, C["dark"], None, False)
    add_shape(slide, MSO_SHAPE.OVAL, x + 21, y + 17, 7, 7, C["red"], None)
    add_shape(slide, MSO_SHAPE.OVAL, x + 35, y + 17, 7, 7, "F59E0B", None)
    add_shape(slide, MSO_SHAPE.OVAL, x + 49, y + 17, 7, 7, "22C55E", None)
    add_picture_fit(slide, path, x + 8, y + 40, w - 16, h - 48)


def class_box(slide, name, attrs, x, y, w=180, h=98, accent=C["blue"]):
    add_rect(slide, x, y, w, h, C["white"], C["border"])
    add_rect(slide, x, y, w, 27, C["navy"], None, False)
    add_rect(slide, x, y, 5, h, accent, None, False)
    text_box(slide, "ENTITY", x + 11, y + 5, 43, 8, 5, "B0DCEB")
    text_box(slide, name, x + 56, y + 6, w - 64, 13, 8, C["white"], True)
    add_line(slide, x + 12, y + 34, x + w - 12, y + 34, "E2ECF7", 0.4)
    text_box(slide, "\n".join(attrs), x + 14, y + 42, w - 28, h - 46, 7, C["ink"])


def flow(slide, steps, x=50, y=184, w=860):
    gap = 18
    bw = (w - gap * (len(steps) - 1)) / len(steps)
    for i, step in enumerate(steps):
        bx = x + i * (bw + gap)
        add_rect(slide, bx, y, bw, 66, C["white"], C["border"])
        add_shape(slide, MSO_SHAPE.OVAL, bx + 10, y + 14, 25, 25, C["blue"], None)
        text_box(slide, str(i + 1), bx + 18, y + 20, 8, 8, 7, C["white"], True)
        text_box(slide, step, bx + 42, y + 17, bw - 52, 28, 9, C["navy"], True)
        if i < len(steps) - 1:
            add_line(slide, bx + bw, y + 33, bx + bw + gap, y + 33, C["cyan"], 1.5)


def new_slide(prs, section, title=None, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, section, len(prs.slides))
    if title:
        section_title(slide, section, title, subtitle or "")
    return slide


def screenshot_slide(prs, title, file_name, notes, label="DEMONSTRATION V3"):
    s = new_slide(prs, label, title, "Capture réelle de l'application et rôle du module dans le workflow global.")
    screenshot(s, SHOT_DIR / file_name, 46, 155, 650, 335)
    add_rect(s, 724, 155, 184, 335, C["white"], C["border"])
    add_rect(s, 724, 155, 184, 5, C["cyan"], None, False)
    text_box(s, "À présenter", 744, 178, 140, 22, 15, C["blue"], True)
    bullet_box(s, notes[:4], 744, 218, 140, 190, 10)
    text_box(s, "Message clé", 744, 438, 140, 16, 9, C["muted"], True)
    text_box(s, notes[-1] if notes else "Module intégré dans V3.", 744, 458, 140, 28, 9, C["ink"])


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    # remove default first slide if any? Presentation starts empty enough via slide layouts.

    s = new_slide(prs, "PFE")
    # Cover grid: 56px outer margins, 430px identity panel, 44px gutter,
    # 374px visual column. Both columns share the same visual center.
    left_x, left_y, left_w, left_h = 56, 74, 430, 414
    right_x, right_w = 530, 374
    add_rect(s, left_x, left_y, left_w, left_h, C["white"], C["border"])
    if LOGO.exists():
        add_picture_fit(s, LOGO, left_x + 34, left_y + 24, 176, 64)
    text_box(s, "PROJET DE FIN D'ÉTUDES", left_x + 34, left_y + 122, 280, 18, 10, C["cyan"], True)
    text_box(s, "Focal V3", left_x + 34, left_y + 160, 370, 46, 33, C["navy"], True)
    text_box(
        s,
        "Plateforme full stack d'assistance opérationnelle, knowledge management, IA et pilotage support.",
        left_x + 34,
        left_y + 222,
        362,
        74,
        15,
        C["ink"],
        True,
        tight=False,
    )
    text_box(s, "Cycle d'ingénieur | ISMAGI | 2025 - 2026", left_x + 34, left_y + 318, 360, 22, 12.5, C["muted"])
    text_box(s, "Abdelfattah AZELMADI", left_x + 34, left_y + 352, 360, 24, 15.5, C["navy"], True)
    add_rect(s, left_x + 34, left_y + 390, 178, 28, C["navy"], C["cyan"])
    text_box(s, "Soutenance PFE", left_x + 58, left_y + 397, 130, 11, 9, C["white"], True, "center")
    screenshot(s, SHOT_DIR / "02-command-center.png", right_x, 112, right_w, 248)
    add_card(s, "Positionnement", "Un cockpit unique pour agents, superviseurs, QA, admins et opérations.", right_x, 388, right_w, 100, C["cyan"])

    s = new_slide(prs, "SYNTHÈSE", "Résumé exécutif", "Le projet répond à un besoin réel : rendre les opérations support plus rapides, cohérentes et mesurables.")
    add_card(s, "Problème", "Outils dispersés, procédures non centralisées, faible traçabilité.", 60, 180, 255, 125, C["red"])
    add_card(s, "Solution", "V3 Shell : CRM, KB, IA, chat, QA, formation, rôles et escalation.", 352, 180, 255, 125, C["blue"])
    add_card(s, "Valeur", "Productivité agent, qualité de réponse, pilotage et scalabilité.", 644, 180, 255, 125, C["green"])
    panel(s, "Angle de soutenance", [
        "Présenter Focal V3 comme une plateforme produit, pas comme une simple application.",
        "Montrer la maîtrise full stack : Angular, Spring Boot, MySQL, sécurité, WebSocket, Docker.",
        "Mettre en avant la valeur métier : support fintech, knowledge, quality, automation.",
    ], 96, 345, 770, 125, C["cyan"])

    s = new_slide(prs, "PLAN", "Sommaire de soutenance", "Structure académique complète, enrichie par des diagrammes techniques et des captures réelles.")
    flow(s, ["Contexte", "Besoin", "Solution", "Conception", "Réalisation", "Démo", "Bilan"], 45, 178, 870)
    panel(s, "Fil conducteur", [
        "Partir du problème opérationnel d'une fintech.",
        "Démontrer l'architecture et les choix de conception.",
        "Prouver le résultat avec des captures réelles de chaque page V3.",
        "Conclure sur les limites, perspectives et valeur professionnelle.",
    ], 90, 322, 780, 145, C["blue"])

    core = [
        ("CONTEXTE", "Contexte du projet", "Les fintechs doivent traiter vite, juste et de manière traçable.", [
            "Le support client est critique dans une fintech : il influence confiance, rétention et conformité.",
            "Les agents manipulent des cas sensibles : KYC, paiement, remboursement, blocage compte, réclamation.",
            "La connaissance doit être centralisée, mesurable et directement exploitable pendant le traitement.",
        ]),
        ("PROBLÉMATIQUE", "Problématique", "Comment industrialiser le support client sans sacrifier la qualité ?", [
            "Centraliser les interactions, procédures, formations et escalades dans un même cockpit.",
            "Garantir la cohérence des réponses entre agents et équipes.",
            "Rendre les opérations mesurables : temps article, statuts, QA, tickets, usage IA.",
            "Sécuriser les données et limiter les modules selon le rôle.",
        ]),
        ("OBJECTIFS", "Objectifs", "Construire un outil opérationnel, pas seulement une interface.", [
            "Améliorer la productivité agents avec CRM, KB, SOP maps et Magic Assistance.",
            "Donner aux managers des dashboards exploitables : KB Analytics, adherence, QA.",
            "Ajouter une IA contrôlée via prompts par rôle et conversations historisées.",
            "Préparer le produit à un déploiement Docker, cloud ou self-hosted.",
        ]),
        ("ANALYSE", "Analyse des besoins", "Les besoins sont structurés par acteur et par responsabilité.", [
            "Agent : traiter, chercher, suivre une procédure, escalader et apprendre.",
            "Team Leader : suivre son équipe, analyser les statuts et les performances.",
            "QA : évaluer les cas, historiser les scores et donner du feedback.",
            "Admin/OPS : gérer comptes, droits, workflows, articles, prompts et intégrations.",
        ]),
        ("CAHIER DES CHARGES", "Cahier des charges", "Les exigences couvrent la fonction, la sécurité et l'exploitation.", [
            "Fonctionnel : CRM, KB, chat, IA, training, QA, roles, escalation, builders.",
            "Technique : API REST, WebSocket, base relationnelle, Docker, intégration Gmail/Jira/OpenAI.",
            "Sécurité : JWT, comptes actifs/désactivés, RoleGuard et contrôle backend.",
            "UX : shell stable, navigation lisible, maps visuelles, feedback immédiat.",
        ]),
        ("EXISTANT", "Étude de l'existant", "Le marché est mature, mais fragmenté selon les usages.", [
            "Zendesk est fort en support, mais coûteux et parfois lourd pour les workflows internes.",
            "Helpjuice est fort en Knowledge Base, mais moins centré sur CRM, QA et opérations.",
            "Slack est excellent en collaboration, mais pas relié nativement aux SOP et analytics KB.",
            "Jira est puissant pour L2/tech, mais pas adapté seul au travail quotidien L1.",
        ]),
        ("SOLUTION", "Solution proposée", "Focal V3 regroupe support, connaissance, IA et pilotage.", [
            "Un shell V3 unique avec sidebar, topbar et routes role-based.",
            "Des modules indépendants mais cohérents : CRM, KB, Magic Assistance, Academy, QA, Chat.",
            "Des builders visuels pour adapter les processus sans modification du code.",
            "Une architecture extensible vers SaaS, self-hosted, SSO et canaux omnicanaux.",
        ]),
    ]
    for label, title, subtitle, bullets in core:
        s = new_slide(prs, label, title, subtitle)
        panel(s, "Points essentiels", bullets, 78, 170, 805, 240, C["blue"])

    s = new_slide(prs, "TECHNOLOGIES", "Stack technique utilisée", "Une architecture moderne, lisible et cohérente avec un projet full stack professionnel.")
    flow(s, ["Angular 21", "Spring Boot 4", "REST", "WebSocket", "MySQL", "Docker"], 42, 178, 875)
    panel(s, "Choix techniques", [
        "Angular : composants standalone, routes lazy-loaded, services HTTP, UI modulaire.",
        "Spring Boot : controllers, services, repositories JPA, security filters.",
        "MySQL/MariaDB : persistence métier, historiques, logs, relations et analytics.",
        "Docker : packaging, isolation et préparation au déploiement cloud/VPS.",
    ], 78, 322, 805, 145, C["cyan"])

    s = new_slide(prs, "CONCEPTION", "Conception UX/UI", "Une interface structurée pour réduire la charge cognitive et accélérer les actions quotidiennes.")
    add_card(s, "Shell stable", "Sidebar + topbar constantes pour garder les repères.", 62, 180, 250, 115, C["blue"])
    add_card(s, "Maps visuelles", "Canvas, nodes, links, zoom, pan et inspecteurs.", 355, 180, 250, 115, C["cyan"])
    add_card(s, "Feedback", "Toasts, loaders, status, messages, erreurs et confirmations.", 648, 180, 250, 115, C["green"])
    screenshot(s, SHOT_DIR / "27-role-access-map.png", 120, 328, 720, 150)

    s = new_slide(prs, "DIAGRAMME DE CAS D'UTILISATION", "Vue globale des acteurs", "Chaque acteur a un périmètre fonctionnel précis, contrôlé par rôle et feature access.")
    flow(s, ["Client public", "Agent", "Team Leader", "QA", "Admin", "OPS"], 42, 178, 875)
    panel(s, "Cas d'utilisation majeurs", [
        "Client public : démarrer une conversation, suivre le bot, escalader, donner un CSAT.",
        "Agent : traiter CRM, consulter KB, suivre Magic Assistance, collaborer, se former.",
        "Supervision : analyser adherence, QA, KB usage, tickets et activité équipe.",
        "Administration : gérer comptes, rôles, maps, articles, prompts, chat projects.",
    ], 78, 322, 805, 145, C["cyan"])

    s = new_slide(prs, "CAS D'UTILISATION", "Détail par responsabilité", "Lecture plus précise des responsabilités métier dans l'application.")
    add_card(s, "Agent", "CRM, KB, SOP, Magic Assistance, formation, chat assigné.", 62, 180, 250, 130, C["blue"])
    add_card(s, "Supervision", "Adherence, QA, analytics, tickets équipe, suivi performance.", 355, 180, 250, 130, C["cyan"])
    add_card(s, "Admin/OPS", "Comptes, rôles, builders, prompts, articles, intégrations.", 648, 180, 250, 130, C["green"])
    panel(s, "Principe", [
        "Chaque action visible doit correspondre à un droit.",
        "Les modules non reliés au rôle ne doivent pas apparaître.",
        "La sécurité doit être appliquée côté backend, pas uniquement côté UI.",
    ], 120, 350, 720, 125, C["blue"])

    s = new_slide(prs, "DIAGRAMME DE CLASSES", "Vue d'ensemble du modèle", "Le modèle est découpé par bounded contexts pour garder un diagramme lisible.")
    flow(s, ["Identity", "Knowledge", "CRM", "Collab", "Chat", "AI", "Training", "Ops"], 35, 178, 890)
    panel(s, "Lecture du modèle", [
        "User est l'entité transversale pour rôles, settings, présence, QA et assignations.",
        "KbArticle relie contenu, catégories, maps, feedbacks et tracking temps.",
        "ChatProject isole workflows, queues, agents, API keys et sessions publiques.",
        "EscalationTicket sépare le support L1/L2 du CRM email et du suivi Jira.",
    ], 78, 322, 805, 145, C["cyan"])

    class_slides = [
        ("Identity & Access", [
            ("User", ["id", "email", "role", "active", "status"], 58, 208, 190, 110, C["blue"]),
            ("UserStatus", ["userId", "status", "updatedAt"], 285, 208, 180, 100, C["cyan"]),
            ("UserStatusHistory", ["fromStatus", "toStatus", "createdAt"], 505, 208, 195, 100, C["cyan"]),
            ("RoleAccessProfile", ["roleKey", "label", "active"], 740, 208, 170, 100, C["green"]),
            ("RoleAccessLink", ["roleKey", "featureKey", "enabled"], 170, 354, 190, 100, C["green"]),
            ("TeamLink", ["agentId", "tlId", "qaId"], 405, 354, 180, 100, C["orange"]),
            ("UserSettings", ["language", "timezone", "jira", "notifications"], 630, 354, 195, 100, C["blue"]),
        ]),
        ("Knowledge & Decision", [
            ("KbCategory", ["id", "name", "order"], 45, 204, 170, 96, C["blue"]),
            ("KbArticle", ["title", "contentHtml", "active"], 255, 204, 185, 105, C["blue"]),
            ("KbMapNode", ["articleId", "title", "x", "y"], 480, 204, 190, 105, C["cyan"]),
            ("KbMapEdge", ["source", "target", "label"], 710, 204, 185, 96, C["cyan"]),
            ("KbTimeLog", ["articleId", "userId", "seconds"], 45, 354, 185, 96, C["green"]),
            ("Page", ["title", "content", "tag"], 280, 354, 170, 90, C["orange"]),
            ("Choice", ["label", "sourcePage", "targetPage"], 500, 354, 190, 96, C["orange"]),
            ("CaseTagNode/Edge", ["label", "parent", "target"], 730, 354, 170, 90, C["green"]),
        ]),
        ("CRM & Communication", [
            ("CemContact", ["name", "email", "phone"], 55, 204, 170, 96, C["blue"]),
            ("CemConversation", ["subject", "status", "assignee"], 270, 204, 200, 105, C["blue"]),
            ("CemInternalNote", ["authorId", "content", "createdAt"], 515, 204, 200, 105, C["cyan"]),
            ("GmailAccount", ["email", "tokens", "expiresAt"], 745, 204, 170, 105, C["green"]),
            ("ChatChannelLink", ["provider", "externalId", "active"], 55, 354, 180, 96, C["orange"]),
            ("ChatConversation", ["channel", "customer", "status"], 285, 354, 180, 96, C["orange"]),
            ("ChatMessage", ["sender", "body", "sentAt"], 515, 354, 190, 100, C["orange"]),
            ("ChatParticipant", ["userId", "role", "joinedAt"], 745, 354, 170, 96, C["orange"]),
        ]),
        ("Live Chat Routing", [
            ("ChatProject", ["name", "slug", "status", "api"], 55, 204, 180, 100, C["cyan"]),
            ("ChatApiKey", ["projectId", "tokenHash", "scope"], 280, 204, 190, 100, C["cyan"]),
            ("ChatWorkflow", ["projectId", "status", "publishedAt"], 515, 204, 190, 96, C["cyan"]),
            ("WorkflowNode", ["type", "payload", "x", "y"], 745, 204, 170, 100, C["cyan"]),
            ("ChatQueue", ["projectId", "name", "status"], 55, 354, 180, 96, C["green"]),
            ("QueueAgent", ["queueId", "userId", "priority"], 280, 354, 190, 96, C["green"]),
            ("LiveSession", ["projectId", "queueId", "status"], 515, 354, 190, 100, C["blue"]),
            ("LiveMessage", ["sessionId", "sender", "body"], 745, 354, 170, 100, C["blue"]),
        ]),
        ("AI, Training, QA & Ops", [
            ("AIConversation", ["userId", "title", "profileId"], 45, 204, 190, 96, C["blue"]),
            ("AIMessage", ["role", "content", "createdAt"], 275, 204, 190, 96, C["blue"]),
            ("PromptProfile", ["prompt", "model", "apiKey"], 505, 204, 190, 96, C["cyan"]),
            ("TrainingAsset", ["type", "path", "duration"], 735, 204, 175, 96, C["green"]),
            ("QaEvaluation", ["ticketId", "agentId", "score"], 160, 354, 190, 96, C["green"]),
            ("EscalationTicket", ["clientId", "priority", "status", "jiraKey"], 395, 354, 205, 100, C["orange"]),
            ("EscalationComment", ["ticketId", "authorId", "body"], 650, 354, 180, 96, C["orange"]),
        ]),
    ]
    for title, boxes in class_slides:
        s = new_slide(prs, "DIAGRAMME DE CLASSES", title, "Les classes sont présentées par domaine pour éviter un schéma illisible.")
        for box in boxes:
            class_box(s, *box)

    inventories = [
        ("Inventaire des classes - Socle", ["User, UserStatus, UserStatusHistory, TeamLink", "RoleAccessProfile, RoleAccessLink, UserSettings", "KbCategory, KbArticle, KbMapNode, KbMapEdge", "KbArticleTimeLog, Page, Choice, CaseTagNode, CaseTagEdge"]),
        ("Inventaire des classes - Communication", ["CemContact, CemConversation, CemInternalNote, GmailAccount", "CollaborationRoom, CollaborationRoomMember, CollaborationMessage", "CollaborationMessageMention, Reaction, ArticleLink", "ChatChannelLink, ChatConversation, ChatMessage, ChatParticipant"]),
        ("Inventaire des classes - Automation & Ops", ["ChatProject, ChatProjectApiKey, ChatWorkflow, ChatWorkflowNode, ChatWorkflowConnection", "ChatQueue, ChatQueueAgent, LiveChatSession, LiveChatMessage, LiveChatAssignment", "ProcessAssistantConversation, ProcessAssistantMessage, PromptProfile, PromptRoleLink", "TrainingAsset, QaEvaluation, EscalationTicket, EscalationComment, WorkflowExecution"]),
    ]
    for title, items in inventories:
        s = new_slide(prs, "INVENTAIRE", title, "Résumé des principales entités persistées dans la base de données.")
        panel(s, "Classes principales", items, 78, 175, 805, 230, C["cyan"])

    flows = [
        ("Login et accès V3", ["Utilisateur", "Login", "Auth API", "JWT", "V3 Shell", "RoleGuard"], ["Saisie credentials", "Validation Spring Security", "Émission JWT", "Chargement profil et droits", "Blocage des pages non autorisées"]),
        ("Lecture article et analytics", ["Agent", "KB", "Article API", "Timer", "Analytics", "Dashboard"], ["Sélection article", "Chargement contenu et SOP map", "Heartbeat temps de lecture", "Feedback like/dislike", "Consolidation par agent et article"]),
        ("Public Chat vers agent", ["Client", "Public Chat", "Workflow", "Queue", "Assign", "Agent"], ["Session publique sans login", "Exécution du decision tree", "Escalade vers queue", "Assignation agent disponible", "Conversation WebSocket + CSAT"]),
        ("Escalade L1/L2 et Jira", ["L1", "Ticket API", "Queue L2", "Playlist", "Jira", "Status"], ["Création ticket avec Client ID", "Détection doublons", "Play verrouille le plus ancien ticket", "Création/suivi Jira via backend", "Résolution contrôlée"]),
    ]
    for title, steps, notes in flows:
        s = new_slide(prs, "SÉQUENCE", title, "Séquence technique de bout en bout, simplifiée pour soutenance.")
        flow(s, steps, 48, 180, 865)
        panel(s, "Lecture technique", notes, 78, 318, 805, 150, C["cyan"])

    arch = [
        ("Modèle de données", ["User centralise rôles, settings, présence et assignations.", "KbArticle relie catégorie, contenu, SOP map, feedback et tracking.", "ChatProject isole workflow, queues, API keys et sessions.", "EscalationTicket relie support L1/L2 et Jira sans mélanger le CRM."]),
        ("Architecture du système", ["Angular V3 gère UI, routes, guards et services HTTP.", "Spring Boot applique sécurité, logique métier, REST et WebSocket.", "MySQL/MariaDB persiste les données relationnelles et historiques.", "Les intégrations externes passent par backend pour protéger les secrets."]),
        ("Architecture de déploiement", ["HTTPS et reverse proxy pour TLS, /api et /ws.", "Frontend Angular build web selon cible.", "Backend Spring Boot en service séparé.", "MySQL avec volume, backups et restrictions réseau."]),
        ("Carte des APIs backend", ["/auth/login : authentification.", "/api/kb, /api/kb-map, /api/kb/analytics : connaissance.", "/api/crm, /api/gmail, /api/case-tags : CRM.", "/api/chat-projects, /api/public/chat, /api/live-chat : chat.", "/api/process-assistant, /api/access-control, /api/escalations : IA, rôles, L1/L2."]),
        ("Développement / Implémentation", ["Frontend : composants Angular standalone et routes lazy-loaded.", "Backend : controllers REST, services métier, repositories JPA.", "Temps réel : WebSocket pour collaboration et live chat.", "Documentation : rapport, guides de déploiement, scripts de génération."]),
    ]
    for title, items in arch:
        s = new_slide(prs, "ARCHITECTURE", title, "Vue technique structurée et adaptée à une présentation d'ingénierie.")
        if "système" in title.lower():
            flow(s, ["Angular V3", "REST API", "Spring Services", "JPA", "MySQL"], 58, 185, 845)
            panel(s, "Responsabilités", items, 100, 325, 760, 140, C["cyan"])
        elif "déploiement" in title.lower():
            flow(s, ["HTTPS", "Reverse Proxy", "Angular", "Spring API", "MySQL", "Volumes"], 40, 185, 880)
            panel(s, "Cible cloud/VPS", items, 100, 325, 760, 140, C["blue"])
        else:
            panel(s, "Détails", items, 78, 175, 805, 245, C["cyan"])

    s = new_slide(prs, "FONCTIONNALITÉS", "Fonctionnalités principales", "Vue claire des modules qui seront démontrés dans la suite.")
    add_card(s, "CRM", "Emails, tickets, notes internes, tags, pièces jointes.", 58, 175, 250, 105, C["blue"])
    add_card(s, "Knowledge", "Articles, SOP maps, feedback, analytics d'usage.", 355, 175, 250, 105, C["cyan"])
    add_card(s, "AI", "Process Copilot, prompts par rôle, historique.", 652, 175, 250, 105, C["green"])
    add_card(s, "Operations", "QA, adherence, training, escalade L1/L2, Jira.", 208, 330, 250, 105, C["orange"])
    add_card(s, "Chat", "Collaboration, live chat, chat projects, workflow bot.", 505, 330, 250, 105, C["blue"])

    screens = [
        ("Login", "01-login.png", ["Authentification sécurisée", "Message d'erreur clair", "Animation login", "Entrée vers V3"]),
        ("Command Center", "02-command-center.png", ["Cockpit central", "Quick jumps", "Navigation par rôle", "Point d'entrée produit"]),
        ("Calendar", "03-calendar.png", ["Planification", "Événements", "Organisation équipe", "Vision temporelle"]),
        ("CRM Inbox", "04-crm-inbox.png", ["Emails et tickets", "Files de traitement", "Statuts", "Travail agent"]),
        ("CRM Ticket Detail", "05-crm-ticket-detail.png", ["Conversation client", "Composer", "Notes internes", "Traçabilité"]),
        ("Magic Assistance", "06-magic-assistance.png", ["Decision tree", "Guidage agent", "Processus", "Cohérence réponse"]),
        ("Knowledge Base", "07-knowledge-base.png", ["Articles", "Recherche", "SOP map", "Feedback agent"]),
        ("KB Analytics", "08-kb-analytics.png", ["Temps de lecture", "Articles consultés", "Agents", "Pilotage knowledge"]),
        ("Collaboration Hub", "09-collaboration-hub.png", ["Canaux", "DM", "Messages", "Coordination interne"]),
        ("Chat Projects", "10-chat-projects.png", ["Chatbots publics", "Workflow", "Queues", "API project"]),
        ("Live Chat", "11-live-chat.png", ["Sessions clients", "Assignation", "Conversation", "Temps réel"]),
        ("Process Copilot", "12-process-assistant.png", ["Assistant IA", "Prompts par rôle", "Historique", "Support agent"]),
        ("Academy - My Trainings", "13-academy-home.png", ["Formations", "Progression", "Parcours", "Montée compétence"]),
        ("Academy - Catalog", "14-academy-catalog.png", ["Catalogue", "Recherche", "Cours", "Autoformation"]),
        ("Academy - Course Viewer", "15-academy-course-viewer.png", ["Lecture cours", "Étapes", "Vidéo", "Suivi temps"]),
        ("Academy - Certificate", "16-academy-certificate.png", ["Certificat", "Validation", "Reconnaissance", "Fin parcours"]),
        ("Academy - Studio", "17-academy-studio.png", ["Création cours", "Modules", "Contenu", "Publication"]),
        ("Academy - Course Editor", "18-academy-course-editor.png", ["Édition", "Sections", "Médias", "Quiz"]),
        ("Academy - Analytics", "19-academy-analytics.png", ["Suivi formation", "Temps", "Completion", "Management"]),
        ("Staff Management", "20-staff-management.png", ["Comptes", "Activation", "Rôles", "Administration"]),
        ("Tree Builder", "21-tree-builder.png", ["Nœuds", "Connecteurs", "Auto-layout", "Save DB"]),
        ("KB Map Builder", "22-kb-map-builder.png", ["SOP article", "Canvas", "Détails", "Pan/zoom"]),
        ("Case Tag Builder", "23-case-tag-builder.png", ["Typologies", "Catégories", "Sous-catégories", "Classification"]),
        ("Prompt Map Builder", "24-prompt-map-builder.png", ["Prompts", "Rôles", "Clés API", "LLM governance"]),
        ("Article Management", "25-article-management.png", ["CRUD articles", "Catégories", "Édition", "Publication"]),
        ("Article Editor", "26-article-editor.png", ["Édition full page", "Contenu riche", "Sauvegarde", "Gestion KB"]),
        ("Role Access Map", "27-role-access-map.png", ["Rôles", "Features", "Liens", "Permissions"]),
        ("Settings", "28-settings.png", ["Profil", "Mot de passe", "Jira", "Préférences"]),
        ("Adherence Dashboard", "29-adherence-dashboard.png", ["Timeline", "Statuts", "Agents", "Supervision"]),
        ("QA Evaluation", "30-qa-evaluation.png", ["Scores", "Évaluation", "Feedback", "Qualité"]),
        ("Team Management", "31-team-management.png", ["TL/QA", "Agents", "Mapping", "Organisation"]),
        ("Channel Access Map", "32-channel-access-map.png", ["Canaux", "Invitations", "Accès", "Collaboration"]),
        ("FlowDesk Board", "33-flowdesk-board.png", ["Board", "Delivery", "Tickets", "Suivi"]),
        ("FlowDesk Backlog", "34-flowdesk-backlog.png", ["Backlog", "Priorités", "Demandes", "Préparation"]),
        ("FlowDesk Sprints", "35-flowdesk-sprints.png", ["Sprints", "Planification", "Exécution", "Suivi"]),
        ("FlowDesk Reports", "36-flowdesk-reports.png", ["Rapports", "KPI", "Analyse", "Delivery"]),
        ("FlowDesk Escalations", "37-flowdesk-escalations.png", ["L1/L2", "Playlist", "Jira", "Résolution"]),
        ("404 Page", "38-v3-not-found.png", ["Page erreur", "Robustesse UX", "Retour", "Navigation"]),
        ("Article Not Found", "39-article-not-found.png", ["Article manquant", "UX erreur", "Retour KB", "Contrôle"]),
    ]
    for title, img, notes in screens:
        screenshot_slide(prs, title, img, notes)

    endings = [
        ("Difficultés rencontrées et solutions", ["Change detection Angular : refresh d'état, listeners et guards.", "Intégrations : Gmail token, Jira CORS, OpenAI quota et API keys.", "Builders visuels : connecteurs, pan/zoom, alignement et lisibilité.", "Production : secrets, Docker, reverse proxy, base persistante et sauvegardes."]),
        ("Résultats obtenus", ["Une plateforme V3 démontrable avec plus de 20 modules.", "39 captures réelles intégrées pour appuyer la soutenance.", "Architecture full stack modulaire et extensible.", "Base solide pour prototype SaaS, self-hosted ou PFE avancé."]),
        ("Limites du projet", ["Remplacer ddl-auto update par migrations versionnées.", "Renforcer tests unitaires, intégration, E2E et charge.", "Ajouter observabilité : logs structurés, metrics, alertes.", "Durcir secrets, permissions serveur et procédure production."]),
        ("Perspectives d'amélioration", ["SSO Keycloak et multi-tenancy.", "RAG et recherche sémantique sur la Knowledge Base.", "Canaux officiels : WhatsApp, Instagram, Messenger, Telegram.", "Kubernetes, CI/CD, monitoring, SLA et business hours."]),
        ("Conclusion", ["Focal V3 dépasse une application CRUD classique.", "Le projet combine produit, architecture, UX, sécurité, IA et opérations.", "La valeur : agents plus rapides, réponses plus cohérentes, pilotage plus mesurable."]),
    ]
    for title, items in endings:
        s = new_slide(prs, "BILAN", title, "Synthèse professionnelle pour conclure la soutenance.")
        panel(s, "À retenir", items, 78, 175, 805, 235, C["cyan"])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, 960, 540, C["navy"], None, False)
    add_rect(s, 0, 0, 960, 6, C["cyan"], None, False)
    text_box(s, "Questions ?", 80, 155, 700, 68, 48, C["white"], True)
    text_box(s, "Merci pour votre attention", 84, 235, 600, 30, 20, C["cyan"], True)
    text_box(s, "Projet de Fin d'Études - Focal V3 - Abdelfattah AZELMADI", 84, 420, 760, 24, 13, "D2DCEB")

    return prs


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    backup = OUT.with_name("presentation-pfe-focal-assist-v3-premium-custom-openxml-backup.pptx")
    if OUT.exists() and not backup.exists():
        shutil.copy2(OUT, backup)
    prs = build()
    prs.save(str(OUT))
    shutil.copy2(OUT, OFFICE_COPY)
    print(OUT)
    print(f"slides={len(prs.slides)}")
    print(OFFICE_COPY)
